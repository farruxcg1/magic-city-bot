"""Namunaviy shartnoma shablonini yaratadi: templates/shartnoma.docx

Magic City o'z shablonini bersa, shu faylni o'sha shablon bilan almashtiramiz va ichiga
{{ ... }} o'rinbosarlarini qo'yamiz (ro'yxati README.md da). Bu skript faqat namuna uchun.

    python scripts/make_template.py
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parent.parent / "templates" / "shartnoma.docx"

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21)
sec.top_margin = sec.bottom_margin = Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.5)

st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.1


def p(text="", bold=False, align=None, size=None, after=None, indent=True):
    par = doc.add_paragraph()
    if align == "c":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "r":
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            par.paragraph_format.first_line_indent = Cm(1)
    if after is not None:
        par.paragraph_format.space_after = Pt(after)
    r = par.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    return par


def h(text):
    par = p(text, bold=True, align="c", after=4)
    par.paragraph_format.space_before = Pt(8)
    par.paragraph_format.keep_with_next = True


def borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        b.append(el)
    pr.append(b)


p("NAMUNA SHABLON — Magic City'ning haqiqiy shartnoma matni bilan almashtiriladi", align="c", size=9, after=6)
p("XONADONNI BO‘LIB TO‘LASH SHARTI BILAN", bold=True, align="c", size=13, after=0)
p("OLDI-SOTDI SHARTNOMASI № {{ number }}", bold=True, align="c", size=13, after=10)

t = doc.add_table(rows=1, cols=2)
t.cell(0, 0).paragraphs[0].add_run("{{ city }} sh.")
r = t.cell(0, 1).paragraphs[0]
r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r.add_run("{{ date_long }}")
p("", after=2)

p("{{ company.name }} (keyingi o‘rinlarda «Sotuvchi») nomidan Ustav asosida ish yurituvchi direktor "
  "{{ company.director or '____________________' }} bir tomondan, va fuqaro {{ client.name }} "
  "(pasport {{ client.passport }}, {{ client.passport_issued }} tomonidan berilgan, JShShIR {{ client.pinfl }}, "
  "yashash manzili: {{ client.address }}) (keyingi o‘rinlarda «Xaridor») ikkinchi tomondan, "
  "ushbu shartnomani quyidagilar haqida tuzdilar:")

h("1. SHARTNOMA PREDMETI")
p("1.1. Sotuvchi {{ object }} tarkibidagi ko‘p qavatli turar-joy binosining {{ unit.block }}-blok, "
  "{{ unit.floor }}-qavatidagi {{ unit.id }} raqamli, {{ unit.rooms }} xonali, umumiy maydoni "
  "{{ unit.area }} m² bo‘lgan xonadonni (keyingi o‘rinlarda «Xonadon») Xaridorga sotadi, Xaridor esa uni "
  "qabul qilib, ushbu shartnomada belgilangan narxni to‘laydi.")
p("1.2. Xonadonning yakuniy maydoni texnik inventarizatsiya (kadastr) ma’lumotlari asosida aniqlanadi.")

h("2. NARX VA TO‘LOV TARTIBI")
p("2.1. Xonadonning narxi {{ price }} ({{ price_words }}) so‘mni tashkil etadi.")
p("{% if installment %}2.2. Xaridor shartnoma imzolangan kuni boshlang‘ich to‘lov sifatida {{ down }} "
  "({{ down_words }}) so‘m to‘laydi. Qolgan {{ rest }} so‘m {{ months }} oy davomida bo‘lib to‘lanadi; "
  "bo‘lib to‘lash uchun qolgan summaga {{ markup }}% ustama qo‘shiladi va bo‘lib to‘lanadigan summa "
  "{{ installment_total }} so‘mni tashkil etadi.{% else %}2.2. Xaridor Xonadon narxini shartnoma "
  "imzolangan kuni to‘liq to‘laydi.{% endif %}")
p("2.3. Shartnoma bo‘yicha jami to‘lanadigan summa {{ total }} ({{ total_words }}) so‘m.")
p("{% if installment %}2.4. Oylik to‘lovlar 1-ilovadagi to‘lov grafigiga muvofiq, birinchi to‘lov "
  "{{ first_due }} dan boshlab, har oyning belgilangan sanasidan kechiktirmay amalga oshiriladi. "
  "Oxirgi to‘lov sanasi: {{ last_due }}.{% else %}2.4. To‘lov naqd yoki Sotuvchining hisob raqamiga "
  "o‘tkazish yo‘li bilan amalga oshiriladi.{% endif %}")
p("2.5. Xaridor to‘lovlarni muddatidan oldin amalga oshirishga haqli.")
p("2.6. Shartnoma narxi taraflarning yozma kelishuvisiz o‘zgartirilmaydi.")

h("3. TARAFLARNING HUQUQ VA MAJBURIYATLARI")
p("3.1. Sotuvchi Xonadonni qurilish loyihasiga muvofiq qurib bitkazish va qabul qilish-topshirish "
  "dalolatnomasi asosida Xaridorga topshirish majburiyatini oladi.")
p("3.2. Xaridor to‘lovlarni grafikda belgilangan muddatlarda to‘liq amalga oshirish majburiyatini oladi.")
p("3.3. Xonadonga mulk huquqi Xaridor shartnoma bo‘yicha barcha to‘lovlarni to‘liq amalga oshirganidan so‘ng "
  "qonunchilikda belgilangan tartibda rasmiylashtiriladi.")

h("4. TARAFLARNING JAVOBGARLIGI")
p("4.1. To‘lov muddati kechiktirilganda Xaridor qonunchilik va taraflar kelishuvida belgilangan "
  "tartibda javobgar bo‘ladi.")
p("4.2. To‘lovlar ketma-ket uch oydan ortiq kechiktirilsa, Sotuvchi shartnomani bir tomonlama bekor "
  "qilish huquqiga ega. Bu holda to‘langan mablag‘lar qaytarish tartibi taraflar kelishuvi bilan belgilanadi.")

h("5. YAKUNIY QOIDALAR")
p("5.1. Shartnoma imzolangan kundan boshlab kuchga kiradi va taraflar o‘z majburiyatlarini to‘liq "
  "bajargunga qadar amal qiladi.")
p("5.2. Nizolar muzokara yo‘li bilan, kelishuvga erishilmasa, sud tartibida hal qilinadi.")
p("5.3. Shartnoma bir xil yuridik kuchga ega ikki nusxada tuzildi, har bir tarafga bittadan.")

h("6. TARAFLARNING REKVIZITLARI")
t = doc.add_table(rows=1, cols=2)
borders(t)
left = t.cell(0, 0)
left.paragraphs[0].add_run("SOTUVCHI").bold = True
for line in ["{{ company.name }}", "Manzil: {{ company.address }}", "STIR: {{ company.inn }}",
             "H/r: {{ company.account }}", "Bank: {{ company.bank }}", "MFO: {{ company.mfo }}",
             "Tel: {{ company.phone }}", "", "Direktor ____________ {{ company.director }}"]:
    left.add_paragraph(line)
right = t.cell(0, 1)
right.paragraphs[0].add_run("XARIDOR").bold = True
for line in ["{{ client.name }}", "Pasport: {{ client.passport }}", "Berilgan: {{ client.passport_issued }}",
             "JShShIR: {{ client.pinfl }}", "Manzil: {{ client.address }}", "Tel: {{ client.phone }}", "", "",
             "Imzo ____________"]:
    right.add_paragraph(line)

# 1-ilova: to'lov grafigi
doc.add_page_break()
p("1-ilova", align="r", after=0)
p("№ {{ number }} shartnomaga", align="r", after=10)
p("TO‘LOV GRAFIGI", bold=True, align="c", size=13, after=2)
p("Xonadon {{ unit.id }}, {{ unit.area }} m², jami {{ total }} so‘m", align="c", after=8)
g = doc.add_table(rows=4, cols=3)
borders(g)
g.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, txt in enumerate(["№", "To‘lov sanasi", "Summa, so‘m"]):
    c = g.cell(0, i).paragraphs[0]
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.add_run(txt).bold = True
g.cell(1, 0).paragraphs[0].add_run("{%tr for r in schedule %}")
for i, txt in enumerate(["{{ r.n }}", "{{ r.date }}", "{{ r.amount }}"]):
    c = g.cell(2, i).paragraphs[0]
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER if i < 2 else WD_ALIGN_PARAGRAPH.RIGHT
    c.add_run(txt)
g.cell(3, 0).paragraphs[0].add_run("{%tr endfor %}")
for row in g.rows:
    row.cells[0].width, row.cells[1].width, row.cells[2].width = Cm(4), Cm(5), Cm(5)
p("", after=10)
t = doc.add_table(rows=1, cols=2)
t.cell(0, 0).paragraphs[0].add_run("Sotuvchi ____________")
r = t.cell(0, 1).paragraphs[0]
r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r.add_run("Xaridor ____________")

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print("Yaratildi:", OUT)
